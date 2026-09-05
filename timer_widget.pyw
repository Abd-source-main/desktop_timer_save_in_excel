"""
Transparent desktop timer widget — circular button edition.

- Floats as a single transparent CIRCLE (no window frame).
- HOVER the circle  -> shows the elapsed time next to it.
- Quick CLICK        -> start / stop the timer (a "period").
- PRESS & HOLD drag  -> move the circle anywhere on screen.
- RIGHT-CLICK        -> menu: toggle "Save to Excel", or Exit.
- Stopping (or exiting) records a period; if "Save to Excel" is on it is
  written to time_log.xlsx, one row per date:
      DATE | Total Hours (decimal) | Total (H:MM) | Period 1 | Period 2 | ...

Run with pythonw.exe so no console appears (use the Desktop shortcut).
"""

import os
import time
import datetime
import tkinter as tk

from openpyxl import Workbook, load_workbook

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(SCRIPT_DIR, "time_log.xlsx")

# Any pixel drawn in this exact colour becomes fully transparent (Windows).
TRANSPARENT_KEY = "#010101"

# How far the pointer must move (pixels) before a press counts as a DRAG
# rather than a CLICK.
DRAG_THRESHOLD = 5

# Circle geometry.
CIRCLE = 48            # diameter
PAD = 6               # padding around circle inside its canvas
CANVAS = CIRCLE + PAD * 2
WIN_W = 180            # total window width (leaves room for the hover time)
WIN_H = CANVAS

COL_IDLE = "#2d7d46"      # green
COL_IDLE_HI = "#3a9a58"
COL_RUN = "#b3402f"       # red
COL_RUN_HI = "#cc4a37"


# Sheet layout: A Date | B Total Hours | C Total (H:MM) | D.. periods
HEADERS = ["Date", "Total Hours", "Total (H:MM)", "Periods ->"]
HM_COL = 3
PERIOD_COL = 4


def fmt_duration(seconds):
    seconds = int(round(seconds))
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def fmt_hm(seconds):
    """Hours + minutes, rounded to the nearest minute: 5400 -> "1:30"."""
    minutes = int(round(seconds / 60.0))
    h, m = divmod(minutes, 60)
    return f"{h}:{m:02d}"


class TimerWidget:
    def __init__(self):
        self.running = False
        self.start_epoch = None
        self.elapsed = 0.0

        self._moved = False
        self._hover = False

        self.root = tk.Tk()
        self.root.title("Timer")
        self.root.overrideredirect(True)
        self.root.attributes("-topmost", True)
        self.root.configure(bg=TRANSPARENT_KEY)
        try:
            self.root.attributes("-transparentcolor", TRANSPARENT_KEY)
        except tk.TclError:
            self.root.attributes("-alpha", 0.9)

        self.root.geometry(f"{WIN_W}x{WIN_H}")

        # Time readout (hidden until hover). Sits to the LEFT of the circle.
        self.time_label = tk.Label(
            self.root, text="00:00:00", font=("Segoe UI", 13, "bold"),
            fg="#00e0a0", bg="#22222a", padx=8, pady=4,
        )
        # placed on demand in _show_time()

        # The circle itself, on a transparent canvas so only the oval shows.
        self.canvas = tk.Canvas(
            self.root, width=CANVAS, height=CANVAS,
            bg=TRANSPARENT_KEY, highlightthickness=0, bd=0, cursor="hand2",
        )
        self.canvas.place(relx=1.0, rely=0.5, anchor="e")

        self.oval = self.canvas.create_oval(
            PAD, PAD, PAD + CIRCLE, PAD + CIRCLE,
            fill=COL_IDLE, outline="#0d0d0f", width=2,
        )
        # Glyph in the centre (play / pause).
        self.glyph = self.canvas.create_text(
            CANVAS / 2, CANVAS / 2, text="▶",  # ▶
            font=("Segoe UI", 16, "bold"), fill="white",
        )

        # --- interaction bindings (on the canvas only) ---
        self.canvas.bind("<ButtonPress-1>", self.on_press)
        self.canvas.bind("<B1-Motion>", self.on_motion)
        self.canvas.bind("<ButtonRelease-1>", self.on_release)
        self.canvas.bind("<Enter>", self.on_enter)
        self.canvas.bind("<Leave>", self.on_leave)
        self.canvas.bind("<Button-3>", self.show_menu)

        # Right-click context menu.
        self.save_var = tk.BooleanVar(value=True)
        self.menu = tk.Menu(self.root, tearoff=0)
        self.menu.add_checkbutton(label="Save to Excel", variable=self.save_var)
        self.menu.add_separator()
        self.menu.add_command(label="Exit", command=self.on_close)

        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

        self.root.update_idletasks()
        self.place_bottom_right()
        self.tick()

    # ------------------------------------------------------------- placement
    def place_bottom_right(self):
        sw = self.root.winfo_screenwidth()
        sh = self.root.winfo_screenheight()
        x = sw - WIN_W - 20
        y = sh - WIN_H - 60
        self.root.geometry(f"{WIN_W}x{WIN_H}+{x}+{y}")

    # ------------------------------------------------------- press vs. drag
    def on_press(self, event):
        self._press_x = event.x_root
        self._press_y = event.y_root
        self._win_x = self.root.winfo_x()
        self._win_y = self.root.winfo_y()
        self._moved = False

    def on_motion(self, event):
        dx = event.x_root - self._press_x
        dy = event.y_root - self._press_y
        if not self._moved and (abs(dx) > DRAG_THRESHOLD or abs(dy) > DRAG_THRESHOLD):
            self._moved = True
            self.time_label.place_forget()  # hide time while dragging
        if self._moved:
            self.root.geometry(f"+{self._win_x + dx}+{self._win_y + dy}")

    def on_release(self, event):
        if self._moved:
            self._moved = False
            if self._hover:
                self._show_time()
            return
        # A genuine click -> toggle the timer.
        self.toggle()

    # ------------------------------------------------------------- hover
    def on_enter(self, event):
        self._hover = True
        self._show_time()

    def on_leave(self, event):
        self._hover = False
        self.time_label.place_forget()

    def _show_time(self):
        self.time_label.config(text=fmt_duration(self.elapsed))
        # Sit just left of the circle.
        self.time_label.place(relx=1.0, rely=0.5, anchor="e", x=-(CANVAS + 2))

    # ------------------------------------------------------------- menu
    def show_menu(self, event):
        try:
            self.menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.menu.grab_release()

    # ------------------------------------------------------------- timer
    def toggle(self):
        self.stop() if self.running else self.start()

    def start(self):
        self.running = True
        self.start_epoch = time.time()
        self.elapsed = 0.0
        self.canvas.itemconfig(self.oval, fill=COL_RUN)
        self.canvas.itemconfig(self.glyph, text="⏸")  # ⏸

    def stop(self):
        if not self.running:
            return
        self.elapsed = time.time() - self.start_epoch
        self.running = False
        self.canvas.itemconfig(self.oval, fill=COL_IDLE)
        self.canvas.itemconfig(self.glyph, text="▶")  # ▶

        if self.save_var.get() and self.elapsed >= 1:
            try:
                save_period(self.elapsed)
            except Exception as exc:
                print("Failed to save:", exc)

        self.elapsed = 0.0
        if self._hover:
            self._show_time()

    def tick(self):
        if self.running:
            self.elapsed = time.time() - self.start_epoch
            if self._hover:
                self.time_label.config(text=fmt_duration(self.elapsed))
        self.root.after(200, self.tick)

    # ------------------------------------------------------------- close
    def on_close(self):
        if self.running:
            self.stop()
        self.root.destroy()

    def run(self):
        self.root.mainloop()


# ----------------------------------------------------------------------------
# Excel writing
# ----------------------------------------------------------------------------
def ensure_hm_column(ws):
    """Upgrade a sheet written before the H:MM column existed."""
    if ws.cell(row=1, column=HM_COL).value == HEADERS[HM_COL - 1]:
        return
    ws.insert_cols(HM_COL)
    ws.cell(row=1, column=HM_COL, value=HEADERS[HM_COL - 1])
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value in (None, ""):
            continue
        total = ws.cell(row=row, column=2).value or 0
        ws.cell(row=row, column=HM_COL, value=fmt_hm(float(total) * 3600.0))


def save_period(seconds):
    """Append a period to today's row in time_log.xlsx.

    Layout, one row per date:
        A: Date | B: Total Hours | C: Total (H:MM) | D: Period 1 | E: Period 2 | ...
    """
    today = datetime.date.today().isoformat()
    hours = seconds / 3600.0
    period_str = fmt_duration(seconds)

    if os.path.exists(EXCEL_PATH):
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
        ensure_hm_column(ws)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Time Log"
        ws.append(HEADERS)

    target_row = None
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=1).value
        if cell is not None and str(cell) == today:
            target_row = row
            break

    if target_row is None:
        target_row = ws.max_row + 1
        total_hours = hours
        ws.cell(row=target_row, column=1, value=today)
        ws.cell(row=target_row, column=PERIOD_COL, value=period_str)
    else:
        prev_total = ws.cell(row=target_row, column=2).value or 0
        total_hours = float(prev_total) + hours
        col = PERIOD_COL
        while ws.cell(row=target_row, column=col).value not in (None, ""):
            col += 1
        ws.cell(row=target_row, column=col, value=period_str)

    ws.cell(row=target_row, column=2, value=round(total_hours, 3))
    ws.cell(row=target_row, column=HM_COL, value=fmt_hm(total_hours * 3600.0))

    wb.save(EXCEL_PATH)


if __name__ == "__main__":
    TimerWidget().run()
